#!/usr/bin/env python3
"""High-reliability bulk import pipeline for Odoo contacts.

Features:
- CSV/XLSX input (XLSX optional: requires openpyxl)
- Validation with per-row error reporting
- Deduplication by email/phone (within file + optional remote Odoo)
- dry-run / apply modes
- Safety guard: real apply requires --explicit-apply
- Logical rollback artifact (JSON with created IDs)
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@dataclass
class RowError:
    row_number: int
    message: str
    row: Dict[str, str]


@dataclass
class RowOk:
    row_number: int
    payload: Dict[str, str]


class OdooClient:
    """Simple interface for Odoo operations.

    In production, wire this class to your XML-RPC/JSON-RPC stack.
    For tests, inject a fake client implementing same methods.
    """

    def __init__(self, enabled: bool = False):
        self.enabled = enabled

    def find_existing_contacts(self, emails: Sequence[str], phones: Sequence[str]) -> Dict[str, int]:
        if not self.enabled:
            return {}
        raise NotImplementedError("Connect your Odoo RPC implementation here")

    def create_contacts(self, records: Sequence[Dict[str, str]]) -> List[int]:
        if not self.enabled:
            return []
        raise NotImplementedError("Connect your Odoo RPC implementation here")

    def deactivate_contacts(self, contact_ids: Sequence[int]) -> int:
        if not self.enabled:
            return 0
        raise NotImplementedError("Connect your Odoo RPC implementation here")


def norm_email(v: str) -> str:
    return (v or "").strip().lower()


def norm_phone(v: str) -> str:
    raw = (v or "").strip()
    return re.sub(r"[^\d+]", "", raw)


def read_input(path: Path) -> List[Dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook  # type: ignore
        except Exception as exc:  # pragma: no cover
            raise RuntimeError("XLSX input requires openpyxl. Install: pip install openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for r in rows[1:]:
            row = {headers[i]: ("" if i >= len(r) or r[i] is None else str(r[i])) for i in range(len(headers))}
            out.append(row)
        return out
    raise ValueError(f"Unsupported input format: {path}")


def validate_and_prepare(rows: Iterable[Dict[str, str]]) -> Tuple[List[RowOk], List[RowError]]:
    ok: List[RowOk] = []
    errors: List[RowError] = []

    for idx, row in enumerate(rows, start=2):  # 2 because header is row 1
        name = (row.get("name") or "").strip()
        email = norm_email(row.get("email", ""))
        phone = norm_phone(row.get("phone", ""))
        mobile = norm_phone(row.get("mobile", ""))

        local_errs = []
        if not name:
            local_errs.append("name is required")
        if email and not EMAIL_RE.match(email):
            local_errs.append("email format invalid")
        if not email and not phone and not mobile:
            local_errs.append("at least one of email/phone/mobile is required")

        if local_errs:
            errors.append(RowError(idx, "; ".join(local_errs), row))
            continue

        payload = {
            "name": name,
            "email": email,
            "phone": phone,
            "mobile": mobile,
            "company_name": (row.get("company_name") or "").strip(),
            "street": (row.get("street") or "").strip(),
            "city": (row.get("city") or "").strip(),
            "state": (row.get("state") or "").strip(),
            "zip": (row.get("zip") or "").strip(),
            "country": (row.get("country") or "").strip(),
        }
        ok.append(RowOk(idx, payload))

    return ok, errors


def dedupe_rows(rows: Sequence[RowOk]) -> Tuple[List[RowOk], List[RowError]]:
    seen = set()
    unique: List[RowOk] = []
    dupes: List[RowError] = []

    for row in rows:
        p = row.payload
        keys = []
        if p.get("email"):
            keys.append(("email", p["email"]))
        if p.get("phone"):
            keys.append(("phone", p["phone"]))
        if p.get("mobile"):
            keys.append(("phone", p["mobile"]))

        hit = next((k for k in keys if k in seen), None)
        if hit:
            dupes.append(RowError(row.row_number, f"duplicate in input by {hit[0]}={hit[1]}", p))
            continue

        for k in keys:
            seen.add(k)
        unique.append(row)

    return unique, dupes


def filter_existing(rows: Sequence[RowOk], client: OdooClient) -> Tuple[List[RowOk], List[RowError]]:
    emails = [r.payload["email"] for r in rows if r.payload.get("email")]
    phones = [r.payload["phone"] for r in rows if r.payload.get("phone")] + [
        r.payload["mobile"] for r in rows if r.payload.get("mobile")
    ]

    existing = client.find_existing_contacts(emails=emails, phones=phones)
    if not existing:
        return list(rows), []

    kept: List[RowOk] = []
    skipped: List[RowError] = []
    for r in rows:
        e = r.payload.get("email")
        p = r.payload.get("phone")
        m = r.payload.get("mobile")
        if (e and f"email:{e}" in existing) or (p and f"phone:{p}" in existing) or (m and f"phone:{m}" in existing):
            skipped.append(RowError(r.row_number, "already exists in Odoo (email/phone)", r.payload))
        else:
            kept.append(r)
    return kept, skipped


def write_report(report_path: Path, report: Dict) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")


def run_pipeline(input_path: Path, mode: str, explicit_apply: bool, client: OdooClient, report_dir: Path) -> int:
    if mode == "apply" and not explicit_apply:
        raise SystemExit("Refusing real import: use --explicit-apply together with --mode apply")

    raw_rows = read_input(input_path)
    valid_rows, validation_errors = validate_and_prepare(raw_rows)
    unique_rows, dedupe_errors = dedupe_rows(valid_rows)
    importable_rows, existing_errors = filter_existing(unique_rows, client)

    created_ids: List[int] = []
    if mode == "apply" and importable_rows:
        created_ids = client.create_contacts([r.payload for r in importable_rows])

    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    report = {
        "timestamp_utc": ts,
        "mode": mode,
        "input": str(input_path),
        "summary": {
            "total_rows": len(raw_rows),
            "valid_rows": len(valid_rows),
            "validation_errors": len(validation_errors),
            "duplicates_in_file": len(dedupe_errors),
            "already_in_odoo": len(existing_errors),
            "ready_to_import": len(importable_rows),
            "created": len(created_ids),
        },
        "errors": {
            "validation": [e.__dict__ for e in validation_errors],
            "duplicates": [e.__dict__ for e in dedupe_errors],
            "existing": [e.__dict__ for e in existing_errors],
        },
        "created_ids": created_ids,
    }

    report_path = report_dir / f"bulk_import_report_{ts}.json"
    write_report(report_path, report)

    if created_ids:
        rollback = {
            "created_ids": created_ids,
            "strategy": "logical_rollback_deactivate",
            "note": "Set active=False for created records to rollback logically.",
        }
        write_report(report_dir / f"bulk_import_rollback_{ts}.json", rollback)

    print(json.dumps({"ok": True, "report": str(report_path), "summary": report["summary"]}, ensure_ascii=False))
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Odoo bulk import pipeline with dry-run/apply safeguards")
    p.add_argument("--input", required=True, help="Path to CSV/XLSX file")
    p.add_argument("--mode", choices=["dry-run", "apply"], default="dry-run")
    p.add_argument("--explicit-apply", action="store_true", help="Mandatory safety flag for --mode apply")
    p.add_argument("--report-dir", default="reports/odoo", help="Directory for JSON reports")
    return p


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    client = OdooClient(enabled=False)
    return run_pipeline(
        input_path=Path(args.input),
        mode=args.mode,
        explicit_apply=args.explicit_apply,
        client=client,
        report_dir=Path(args.report_dir),
    )


if __name__ == "__main__":
    sys.exit(main())
